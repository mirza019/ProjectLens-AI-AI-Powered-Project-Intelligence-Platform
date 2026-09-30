from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session
from app.models.entities import Project
from app.analytics.financial import health_status

def project_report_data(db:Session,project_id:str,period:str)->dict:
    p=db.get(Project,project_id)
    if not p: raise ValueError("Project not found")
    exposure=sum(r.exposure for r in p.risks); variance=p.forecast_cost-p.budget_cost
    return {"title":f"{p.name} Monthly Project Review","period":period,"project":p.name,"executive_summary":f"{p.name} is {p.completion:.0f}% complete with expected margin of {p.expected_margin:.1f}%.","financial_status":{"value":p.value,"budget_cost":p.budget_cost,"forecast_cost":p.forecast_cost,"variance":variance,"expected_margin":p.expected_margin},"schedule":{"delay_days":p.schedule_days},"risks":{"open_count":len(p.risks),"exposure":exposure},"management_attention":["Review adverse forecast drivers"] if variance>0 else ["Maintain current cost controls"],"status":"Draft"}

def render_pdf(data:dict,path:Path)->Path:
    path.parent.mkdir(parents=True,exist_ok=True); c=Canvas(str(path),pagesize=A4); w,h=A4
    c.setFillColorRGB(.07,.07,.06); c.rect(0,0,w,h,fill=1,stroke=0); c.setFillColorRGB(.76,.61,.36); c.setFont("Helvetica-Bold",20); c.drawString(48,h-60,"ProjectLens AI")
    c.setFillColorRGB(.94,.92,.88); c.setFont("Helvetica-Bold",16); c.drawString(48,h-95,data["title"]); c.setFont("Helvetica",9); c.setFillColorRGB(.7,.68,.64); c.drawString(48,h-112,f"Reporting period: {data['period']} · Synthetic demonstration data")
    y=h-155
    for title,value in [("Executive Summary",data["executive_summary"]),("Financial Status",f"Forecast cost: EUR {data['financial_status']['forecast_cost']:,.0f} | Variance: EUR {data['financial_status']['variance']:,.0f} | Margin: {data['financial_status']['expected_margin']:.1f}%"),("Schedule Status",f"Current schedule variance: {data['schedule']['delay_days']} days"),("Key Risks",f"{data['risks']['open_count']} open risks with EUR {data['risks']['exposure']:,.0f} exposure"),("Management Attention","; ".join(data["management_attention"]))]:
        c.setFillColorRGB(.76,.61,.36); c.setFont("Helvetica-Bold",11); c.drawString(48,y,title); y-=18; c.setFillColorRGB(.85,.83,.78); c.setFont("Helvetica",9)
        for line in [value[i:i+95] for i in range(0,len(value),95)]: c.drawString(48,y,line); y-=14
        y-=15
    c.save(); return path

def render_xlsx(data:dict,path:Path)->Path:
    path.parent.mkdir(parents=True,exist_ok=True); wb=Workbook(); ws=wb.active; ws.title="Executive Summary"
    navy="0B1F3A"; blue="0B65D8"; pale="EAF2FC"
    ws.merge_cells("A1:D1"); ws["A1"]="ProjectLens AI — Monthly Project Review"; ws["A1"].font=Font(size=18,bold=True,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor=navy); ws["A1"].alignment=Alignment(horizontal="left")
    ws.append(["Project",data["project"],"Period",data["period"]]); ws.append(["Status","Draft","Data classification","Synthetic demonstration data"]); ws.append([])
    ws.append(["Executive summary",data["executive_summary"]]); ws.merge_cells(start_row=5,start_column=2,end_row=5,end_column=4); ws["A5"].font=Font(bold=True,color=navy)
    ws.append([]); ws.append(["Financial metric","Value","Schedule / risk metric","Value"])
    for cell in ws[7]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor=blue)
    fin=data["financial_status"]; ws.append(["Contract value",fin["value"],"Schedule variance (days)",data["schedule"]["delay_days"]]); ws.append(["Budget cost",fin["budget_cost"],"Open risks",data["risks"]["open_count"]]); ws.append(["Forecast cost",fin["forecast_cost"],"Risk exposure",data["risks"]["exposure"]]); ws.append(["Cost variance",fin["variance"],"Expected margin",fin["expected_margin"]])
    for row in range(8,12):
        ws.cell(row,2).number_format='€#,##0.00'; ws.cell(row,4).number_format='#,##0.00'
        if row%2==0:
            for col in range(1,5): ws.cell(row,col).fill=PatternFill("solid",fgColor=pale)
    ws.append([]); ws.append(["Management attention"]); ws["A13"].font=Font(bold=True,color=navy)
    for item in data["management_attention"]: ws.append(["•",item])
    ws.freeze_panes="A7"; ws.auto_filter.ref="A7:D11"
    for col,width in {"A":24,"B":36,"C":28,"D":22}.items(): ws.column_dimensions[col].width=width
    wb.save(path); return path
